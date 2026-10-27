"""Read and write spreadsheet-style data: Excel (.xlsx) and CSV/TSV text files.

Schools normally keep rosters and scores in Excel, so every import accepts a workbook or a
CSV and returns the same thing: rows of plain text. Nothing here touches the database.

Reading is forgiving about how Excel saves files (UTF-8 with or without BOM, UTF-16, Windows
code page, comma / semicolon / tab separators, numbers stored as 8031234567.0, real dates).
Writing neutralises spreadsheet formulas so imported names can never run as formulas.
"""

from __future__ import annotations

import csv
import io
import zipfile
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_ROWS = 20_000
XLSX_SUFFIXES = {".xlsx", ".xlsm"}
TEXT_SUFFIXES = {".csv", ".txt", ".tsv"}
SPREADSHEET_FILTER = "Spreadsheets (*.xlsx *.csv *.txt *.tsv);;Excel workbooks (*.xlsx);;CSV files (*.csv);;All files (*)"


def cell_text(value) -> str:
    """Convert one spreadsheet cell to the text a person would have seen."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time(0, 0) else value.isoformat(sep=" ", timespec="minutes")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.strftime("%H:%M")
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value != value or value in (float("inf"), float("-inf")):
            return ""
        text = format(Decimal(repr(value)), "f")
        return text[:-2] if text.endswith(".0") else text
    return str(value).replace(" ", " ").strip()


def _decode(raw: bytes) -> str:
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        return raw.decode("utf-16")
    try:
        return raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        return raw.decode("cp1252", errors="replace")  # what Excel uses for "CSV" on Windows


def _delimiter(text: str) -> str:
    first = next((line for line in text.splitlines() if line.strip()), "")
    counts = {d: first.count(d) for d in (",", ";", "\t", "|")}
    best = max(counts, key=counts.get)
    return best if counts[best] else ","


def _read_text(raw: bytes) -> list[list[str]]:
    text = _decode(raw)
    rows = []
    try:
        for cells in csv.reader(io.StringIO(text, newline=""), delimiter=_delimiter(text)):
            rows.append([cell.replace(" ", " ").strip() for cell in cells])
            if len(rows) > MAX_ROWS:
                raise ValueError(f"The file has more than {MAX_ROWS:,} rows.")
    except csv.Error as exc:
        raise ValueError(f"The file could not be read as CSV: {exc}") from exc
    return rows


def _read_xlsx(path: Path) -> list[list[str]]:
    try:
        from openpyxl import load_workbook
        from openpyxl.utils.exceptions import InvalidFileException
    except ImportError as exc:  # keeps the rest of the app working if the package is missing
        raise ValueError("Excel support is not installed. Run: pip install openpyxl  (or save the sheet as CSV).") from exc
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except (zipfile.BadZipFile, InvalidFileException, KeyError, OSError) as exc:
        if isinstance(exc, PermissionError):
            raise
        raise ValueError("This is not a valid Excel (.xlsx) file. Open it in Excel and use Save As > Excel Workbook.") from exc
    try:
        for sheet in workbook.worksheets:
            rows = []
            for values in sheet.iter_rows(values_only=True):
                rows.append([cell_text(v) for v in values])
                if len(rows) > MAX_ROWS:
                    raise ValueError(f"The sheet has more than {MAX_ROWS:,} rows.")
            while rows and not any(rows[-1]):
                rows.pop()
            if any(any(r) for r in rows):
                return rows  # first sheet that has data
        return []
    finally:
        workbook.close()


def read_table(path: str | Path, max_bytes: int = MAX_FILE_BYTES) -> list[list[str]]:
    """Return the file as rows of text. Row ``n`` of the result is row ``n`` in Excel (blanks kept)."""
    path = Path(path)
    size = path.stat().st_size
    if size > max_bytes:
        raise ValueError(f"The file is larger than {max_bytes // (1024 * 1024)} MiB.")
    suffix = path.suffix.lower()
    if suffix == ".xls":
        raise ValueError("Old Excel (.xls) files are not supported. Open it in Excel and use Save As > Excel Workbook (.xlsx).")
    with path.open("rb") as handle:
        raw = handle.read()
    if suffix in XLSX_SUFFIXES or (suffix not in TEXT_SUFFIXES and raw[:2] == b"PK"):
        return _read_xlsx(path)
    if raw[:4] == b"\xd0\xcf\x11\xe0":
        raise ValueError("This looks like an old Excel (.xls) file. Open it in Excel and use Save As > Excel Workbook (.xlsx).")
    return _read_text(raw)


# ----------------------------------------------------------------------------- writing

def _safe(value) -> str:
    text = "" if value is None else str(value)
    return text


def write_workbook(path: str | Path, sheets: list[dict]) -> Path:
    """Write ``sheets`` (each ``{name, headers, rows, text_columns, notes}``) to an .xlsx workbook.

    * All cells are stored as text, so phone numbers keep their leading zero and a value that
      begins with ``=`` is data, never a formula.
    * ``notes``: optional list of (label, text) rows for an instructions sheet.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    path = Path(path)
    if path.suffix.lower() != ".xlsx":
        path = path.with_suffix(".xlsx")
    workbook = Workbook()
    workbook.remove(workbook.active)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="102A43")
    for spec in sheets:
        sheet = workbook.create_sheet(title=(spec["name"] or "Sheet")[:31])
        headers = list(spec.get("headers", []))
        if headers:
            sheet.append(headers)
            for cell in sheet[1]:
                cell.font, cell.fill = header_font, header_fill
                cell.alignment = Alignment(vertical="center")
            sheet.freeze_panes = "A2"
        for row in spec.get("rows", []):
            sheet.append([_safe(v) for v in row])
        for row in sheet.iter_rows(min_row=1):
            for cell in row:
                if cell.value is not None:
                    cell.number_format = "@"
                    cell.data_type = "s"       # text, never a formula
        text_columns = set(spec.get("text_columns", ()))
        for index, header in enumerate(headers, start=1):
            letter = get_column_letter(index)
            longest = max([len(str(header))] + [len(str(r[index - 1])) for r in spec.get("rows", []) if index - 1 < len(r)])
            sheet.column_dimensions[letter].width = max(12, min(longest + 3, 48))
            if header in text_columns:
                for row_number in range(2, max(sheet.max_row, 200) + 1):
                    sheet.cell(row=row_number, column=index).number_format = "@"
        for label, text in spec.get("notes", []):
            sheet.append([label, text])
            sheet.cell(row=sheet.max_row, column=1).font = Font(bold=True)
            sheet.cell(row=sheet.max_row, column=2).alignment = Alignment(wrap_text=True, vertical="top")
        if spec.get("notes"):
            sheet.column_dimensions["A"].width = 22
            sheet.column_dimensions["B"].width = 80
    workbook.save(path)
    return path

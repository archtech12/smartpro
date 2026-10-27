"""Excel and CSV imports: encodings, separators, Excel quirks, templates and exports (synthetic data)."""

import datetime

import pytest

pytest.importorskip("openpyxl")

from openpyxl import Workbook, load_workbook  # noqa: E402

from smartschool import tabular
from smartschool.score_import import preview_score_import, write_score_template
from smartschool.student_import import HEADERS, preview_student_import, write_student_template


def _xlsx(path, rows, sheet="Sheet1"):
    workbook = Workbook()
    workbook.active.title = sheet
    for row in rows:
        workbook.active.append(row)
    workbook.save(path)
    return path


def test_student_import_reads_an_excel_workbook_with_friendly_headings_and_real_cells(tmp_path):
    path = _xlsx(tmp_path / "roster.xlsx", [
        ["First Name", "Surname", "Sex", "Phone", "NIN", "DOB", "E-mail", "Notes"],
        ["Amina", "Okafor", "f", 8031234567, 12345678901.0, datetime.datetime(2012, 3, 4), "amina@example.org", "x"],
        [None] * 8,
        ["Musa", "Bello", "Male", "08039999999", "", "", "", ""],
    ])
    preview = preview_student_import(path)
    assert preview["row_count"] == 2
    assert preview["rows"][0] == {
        "first_name": "Amina", "last_name": "Okafor", "gender": "Female", "guardian_phone": "8031234567",
        "nin": "12345678901", "date_of_birth": "2012-03-04", "email": "amina@example.org",
    }
    assert preview["rows"][1]["guardian_phone"] == "08039999999"          # text keeps its leading zero
    assert preview["warnings"] == ["Ignored columns: Notes"]


@pytest.mark.parametrize(("encoding", "text"), [
    ("utf-8-sig", "first_name;last_name\nJosé;Bello\n"),                    # Excel in a comma-decimal region
    ("cp1252", "first_name,last_name\nJosé,Bello\n"),                       # "CSV" saved by Windows Excel
    ("utf-16", "first_name\tlast_name\r\nJosé\tBello\r\n"),                 # "Unicode Text" saved by Excel
    ("utf-8", "First Name,Last Name\r\nJosé,Bello\r\n"),
])
def test_student_import_reads_csv_in_any_excel_encoding_and_separator(tmp_path, encoding, text):
    path = tmp_path / "roster.csv"
    path.write_bytes(text.encode(encoding))
    assert preview_student_import(path)["rows"][0]["first_name"] == "José"


def test_all_problem_rows_are_reported_together_not_one_at_a_time(tmp_path):
    path = _xlsx(tmp_path / "bad.xlsx", [
        ["first_name", "last_name", "email", "date_of_birth"],
        ["A", "", "", ""], ["B", "B", "nope", ""], ["C", "C", "", "03/04/2012"], ["D", "D", "", "2012-02-30"],
    ])
    with pytest.raises(ValueError) as caught:
        preview_student_import(path)
    message = str(caught.value)
    assert message.startswith("4 row(s) need fixing")
    for expected in ("Row 2: last name is required", "Row 3: Email must be a valid", "Row 4: date of birth 03/04/2012", "Row 5: Date of birth must use"):
        assert expected in message


def test_dates_and_scientific_nin_from_excel_text_are_repaired_only_when_unambiguous(tmp_path):
    path = tmp_path / "s.csv"
    path.write_text("first_name,last_name,nin,date_of_birth\nA,A,1.2345678901E+10,25/12/2012\nB,B,,12/25/2012\nC,C,,5 March 2011\n", encoding="utf-8")
    rows = preview_student_import(path)["rows"]
    assert (rows[0]["nin"], rows[0]["date_of_birth"]) == ("12345678901", "2012-12-25")
    assert rows[1]["date_of_birth"] == "2012-12-25" and rows[2]["date_of_birth"] == "2011-03-05"
    path.write_text("first_name,last_name,date_of_birth\nA,A,03/04/2012\n", encoding="utf-8")
    with pytest.raises(ValueError, match="could mean day/month or month/day"):
        preview_student_import(path)


def test_missing_required_columns_error_lists_what_was_found_and_bad_files_are_explained(tmp_path):
    path = _xlsx(tmp_path / "x.xlsx", [["Name", "Class"], ["Amina Okafor", "JSS 1"]])
    with pytest.raises(ValueError, match="missing required columns: first_name, last_name.*Name, Class"):
        preview_student_import(path)
    fake = tmp_path / "fake.xlsx"
    fake.write_text("not a workbook")
    with pytest.raises(ValueError, match="not a valid Excel"):
        preview_student_import(fake)
    old = tmp_path / "old.xls"
    old.write_bytes(b"\xd0\xcf\x11\xe0 old binary")
    with pytest.raises(ValueError, match="Old Excel"):
        preview_student_import(old)
    empty = tmp_path / "empty.csv"
    empty.write_text("")
    with pytest.raises(ValueError, match="empty"):
        preview_student_import(empty)


def test_excel_template_has_headers_text_columns_and_an_instructions_sheet_and_round_trips(tmp_path):
    path = write_student_template(tmp_path / "template.xlsx")
    workbook = load_workbook(path)
    assert workbook.sheetnames == ["Students", "Instructions"]
    assert [c.value for c in workbook["Students"][1]] == list(HEADERS)
    assert workbook["Students"]["D2"].number_format == "@"            # phone column is text
    with pytest.raises(ValueError, match="No student rows"):
        preview_student_import(path)                                   # blank template is rejected clearly
    csv_path = write_student_template(tmp_path / "template.csv")
    assert csv_path.read_text(encoding="utf-8-sig").strip() == ",".join(HEADERS)


def test_workbooks_never_store_formulas_and_exports_keep_text(tmp_path):
    path = tabular.write_workbook(tmp_path / "out.xlsx", [{"name": "S", "headers": ["name", "phone"], "rows": [["=HYPERLINK(1)", "0803"]]}])
    cell = load_workbook(path)["S"]["A2"]
    assert cell.data_type == "s" and cell.value == "=HYPERLINK(1)"
    assert tabular.read_table(path)[1] == ["=HYPERLINK(1)", "0803"]


def test_score_sheet_round_trips_through_excel(tmp_path):
    components = [{"id": 1, "component_name": "CA Score", "max_score": "40"}, {"id": 2, "component_name": "Exam Score", "max_score": "60"}]
    students = [{"admission_no": "A-1", "first_name": "Amina", "last_name": "Okafor", "enrollment_id": 11, "scores": {}},
                {"admission_no": "A-2", "first_name": "Musa", "last_name": "Bello", "enrollment_id": 12, "scores": {}}]
    path = tmp_path / "scores.xlsx"
    write_score_template(path, students, components, "ENG")
    workbook = load_workbook(path)
    sheet = workbook.active
    assert [c.value for c in sheet[1]] == ["admission_no", "first_name", "last_name", "ENG_CA", "ENG_EXAM"]
    # a teacher fills the sheet in Excel with real numbers
    sheet["D2"], sheet["E2"], sheet["D3"] = 35.5, 50, 40
    for cell in ("D2", "E2", "D3"):
        sheet[cell].data_type = "n"
    workbook.save(path)
    result = preview_score_import(path, students, components, "ENG")
    assert result["updates"] == {11: {1: "35.5", 2: "50"}, 12: {1: "40"}}
    with pytest.raises(ValueError, match="exceeds|between|maximum|at most|must"):
        sheet["D3"] = 99
        workbook.save(path)
        preview_score_import(path, students, components, "ENG")


def test_module_export_to_excel_is_admin_free_read_and_keeps_leading_zeros(tmp_path, monkeypatch):
    from smartschool import database as db

    monkeypatch.setattr(db, "APP_DIR", tmp_path)
    monkeypatch.setattr(db, "DB_PATH", tmp_path / "school.db")
    db.initialize()
    db.create_user("admin", "synthetic-hash")
    db.add_row("Students", ["ADM-1", "Amina", "Okafor", "Female", "", "08031234567"], actor="admin")
    out = db.export_module_xlsx("Students", str(tmp_path / "students.xlsx"))
    rows = tabular.read_table(out)
    assert rows[0][:2] == ["admission_no", "first_name"]
    assert "08031234567" in rows[1]
